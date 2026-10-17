"""建立字根拆解資料庫 data/morph.db（build.sh 會自動執行）。

來源：
- MorphoLex-en（約 7 萬字的字首／字根／字尾拆解，CC BY-NC-SA 4.0，限非商業使用）
  https://github.com/hugomailhot/MorphoLex-en
  拆解格式：{<dis<(appear)>ance>}  ← <…< 字首、(…) 字根、>…> 字尾
- ECDICT wordroot.txt（約 600 個字根與字首字尾的意思、來源、例字，MIT）
  https://github.com/skywind3000/ECDICT

wordroot 的英文意思用 Seed-X 翻成中文，結果快取在 data/wordroot_zh.json（只翻一次）。

用法：
    uv run python build_morph.py <MorphoLEX_en.xlsx> <wordroot.txt> <輸出 morph.db>
"""

import json
import re
import sqlite3
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
MODEL = ROOT / "models" / "seedx-mlx-8bit"
ZH_CACHE = ROOT / "data" / "wordroot_zh.json"
# 字首字尾的中文意思（人工校對，覆蓋機器翻譯）；字根仍用 Seed-X 翻譯
AFFIX_ZH = json.loads((Path(__file__).resolve().parent / "affix_zh.json").read_text())


def read_morpholex(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    out = {}
    for name in wb.sheetnames:
        if name == "Presentation" or name.startswith("All"):
            continue
        rows = wb[name].iter_rows(values_only=True)
        header = next(rows)
        try:
            wi, si = header.index("Word"), header.index("MorphoLexSegm")
        except ValueError:
            continue
        for r in rows:
            word, segm = r[wi], r[si]
            if word and segm:
                out.setdefault(str(word).lower(), str(segm))
    return out


def translate_meanings(entries):
    """把 wordroot 的英文意思翻成中文（Seed-X），已翻過的從快取讀。"""
    cache = json.loads(ZH_CACHE.read_text()) if ZH_CACHE.exists() else {}
    todo = sorted({e.get("meaning", "").strip() for e in entries.values()} - set(cache) - {""})
    if todo:
        from mlx_lm import generate, load
        from mlx_lm.sample_utils import make_sampler
        from opencc import OpenCC

        s2twp = OpenCC("s2twp")
        print(f"用 Seed-X 翻譯 {len(todo)} 條字根意思（只需做一次）…", flush=True)
        model, tok = load(str(MODEL))
        sampler = make_sampler(temp=0.0)
        t0 = time.perf_counter()
        for i, m in enumerate(todo, 1):
            prompt = f"Translate the following English sentence into Chinese:\n{m} <zh>"
            zh = generate(model, tok, prompt, max_tokens=80, sampler=sampler)
            zh = re.sub(r"</?s>", "", zh).split("\n")[0].strip()
            cache[m] = s2twp.convert(zh)
            if i % 50 == 0:
                print(f"  {i}/{len(todo)}（{time.perf_counter() - t0:.0f} 秒）", flush=True)
                ZH_CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1))
        ZH_CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1))
    return cache


def main():
    xlsx, wordroot, dst = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
    entries = json.loads(wordroot.read_text())
    zh = translate_meanings(entries)
    words = read_morpholex(xlsx)

    tmp = dst.with_suffix(".tmp")
    tmp.unlink(missing_ok=True)
    conn = sqlite3.connect(tmp)
    conn.execute("CREATE TABLE morph(word TEXT PRIMARY KEY COLLATE NOCASE, segm TEXT NOT NULL)")
    conn.executemany("INSERT INTO morph VALUES (?,?)", words.items())
    conn.execute("""CREATE TABLE affix(key TEXT PRIMARY KEY, class TEXT, meaning_en TEXT, meaning_zh TEXT,
                    origin TEXT, examples TEXT)""")
    conn.executemany("INSERT INTO affix VALUES (?,?,?,?,?,?)", [
        (k, e.get("class", ""), e.get("meaning", ""),
         AFFIX_ZH.get(k) or zh.get(e.get("meaning", "").strip(), ""),
         e.get("origin", ""), json.dumps(e.get("example", []), ensure_ascii=False))
        for k, e in entries.items()
    ])
    conn.commit()
    conn.close()
    tmp.replace(dst)
    print(f"字根資料庫完成：{len(words)} 個字的拆解、{len(entries)} 個字根／字首字尾 → {dst}")


if __name__ == "__main__":
    main()
